import json
import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

base_dir = r"D:\project_model"
json_path = os.path.join(base_dir, "models", "model_comparison.json")
excel_en_path = os.path.join(base_dir, "models_benchmark_evaluation_en.xlsx")

with open(json_path, "r", encoding="utf-8") as f:
    data = json.load(f)

with pd.ExcelWriter(excel_en_path, engine="openpyxl") as writer:
    # ---------------- 1. Models Benchmark Sheet ----------------
    benchmark_data = [
        {
            "Rank": 1,
            "Model (Algorithm)": "Linear SVM (Calibrated)",
            "Accuracy": 0.9501,
            "F1-Score": 0.9536,
            "Precision": 0.9526,
            "Recall": 0.9546,
            "5-Fold CV Mean": 0.9503,
            "Evaluation / Status": "Production Champion (Best Overall Balance & Generalization)"
        },
        {
            "Rank": 2,
            "Model (Algorithm)": "Logistic Regression",
            "Accuracy": 0.9481,
            "F1-Score": 0.9519,
            "Precision": 0.9479,
            "Recall": 0.9559,
            "5-Fold CV Mean": 0.9481,
            "Evaluation / Status": "Strong Competitor (High Accuracy & Stable Probabilities)"
        },
        {
            "Rank": 3,
            "Model (Algorithm)": "Multinomial Naive Bayes",
            "Accuracy": 0.9354,
            "F1-Score": 0.9416,
            "Precision": 0.9155,
            "Recall": 0.9693,
            "5-Fold CV Mean": 0.9407,
            "Evaluation / Status": "Fast Baseline (High Recall, Higher False Positives)"
        },
        {
            "Rank": 4,
            "Model (Algorithm)": "Random Forest",
            "Accuracy": 0.8291,
            "F1-Score": 0.8607,
            "Precision": 0.7658,
            "Recall": 0.9824,
            "5-Fold CV Mean": 0.8624,
            "Evaluation / Status": "Biased towards Majority Class (Overfitting on High-Dimensional Sparse TF-IDF)"
        }
    ]
    df_bm = pd.DataFrame(benchmark_data)
    df_bm.to_excel(writer, sheet_name="Models Benchmark", index=False)

    # ---------------- 2. Train-Test Split & Dataset Specs ----------------
    split_data = [
        {
            "Dataset Split": "Training Set",
            "Percentage": "80.0%",
            "Samples Count": 17720,
            "Sampling Strategy": "Stratified Random Sampling (random_state=42)",
            "Purpose": "Algorithm learning, weight optimization & TF-IDF vocabulary fitting"
        },
        {
            "Dataset Split": "Testing Set (Holdout)",
            "Percentage": "20.0%",
            "Samples Count": 4430,
            "Sampling Strategy": "Stratified Holdout (Unseen Data)",
            "Purpose": "Unbiased final model benchmarking and generalization verification"
        },
        {
            "Dataset Split": "Total Corpus",
            "Percentage": "100.0%",
            "Samples Count": 22150,
            "Sampling Strategy": "Bilingual Augmented Dataset (Arabic + English)",
            "Purpose": "Multi-dialect, multi-category sentiment coverage"
        },
        {
            "Dataset Split": "5-Fold Cross Validation",
            "Percentage": "20% per fold (in Train)",
            "Samples Count": 3544,
            "Sampling Strategy": "5-Fold Stratified K-Fold CV",
            "Purpose": "Ensures no overfitting and stable performance across different folds"
        },
        {
            "Dataset Split": "Positive Class Representation",
            "Percentage": "53.8%",
            "Samples Count": 11917,
            "Sampling Strategy": "Balanced Binary Class Mapping (Scores 4 & 5)",
            "Purpose": "Representative sample of satisfied customers"
        },
        {
            "Dataset Split": "Negative Class Representation",
            "Percentage": "46.2%",
            "Samples Count": 10233,
            "Sampling Strategy": "Balanced Binary Class Mapping (Scores 1 & 2)",
            "Purpose": "Representative sample of complaints and product defect feedback"
        }
    ]
    df_split = pd.DataFrame(split_data)
    df_split.to_excel(writer, sheet_name="Train-Test Split Specs", index=False)

    # ---------------- 3. Confusion Matrix (Winning Model) ----------------
    cm_data = [
        {
            "Actual Class \\ Predicted Class": "Actual Negative (Defect / Unfavorable)",
            "Predicted Negative": 1937,
            "Predicted Positive": 113,
            "Actual Total": 2050,
            "Class Accuracy / Specificity": "94.49%"
        },
        {
            "Actual Class \\ Predicted Class": "Actual Positive (Satisfied / Favorable)",
            "Predicted Negative": 108,
            "Predicted Positive": 2272,
            "Actual Total": 2380,
            "Class Accuracy / Sensitivity": "95.46%"
        },
        {
            "Actual Class \\ Predicted Class": "Total Predicted",
            "Predicted Negative": 2045,
            "Predicted Positive": 2385,
            "Actual Total": 4430,
            "Class Accuracy / Specificity": "Overall Acc: 95.01%"
        }
    ]
    df_cm = pd.DataFrame(cm_data)
    df_cm.to_excel(writer, sheet_name="Confusion Matrix (Best Model)", index=False)

    # ---------------- 4. All Models Confusion Comparison ----------------
    all_cm = [
        {
            "Model Name": "Linear SVM (Calibrated)",
            "True Negative (TN)": 1937,
            "False Positive (FP)": 113,
            "False Negative (FN)": 108,
            "True Positive (TP)": 2272,
            "Total Test Samples": 4430,
            "Precision": "95.26%",
            "Recall": "95.46%",
            "F1-Score": "95.36%"
        },
        {
            "Model Name": "Logistic Regression",
            "True Negative (TN)": 1925,
            "False Positive (FP)": 125,
            "False Negative (FN)": 105,
            "True Positive (TP)": 2275,
            "Total Test Samples": 4430,
            "Precision": "94.79%",
            "Recall": "95.59%",
            "F1-Score": "95.19%"
        },
        {
            "Model Name": "Multinomial Naive Bayes",
            "True Negative (TN)": 1837,
            "False Positive (FP)": 213,
            "False Negative (FN)": 73,
            "True Positive (TP)": 2307,
            "Total Test Samples": 4430,
            "Precision": "91.55%",
            "Recall": "96.93%",
            "F1-Score": "94.16%"
        },
        {
            "Model Name": "Random Forest",
            "True Negative (TN)": 1335,
            "False Positive (FP)": 715,
            "False Negative (FN)": 42,
            "True Positive (TP)": 2338,
            "Total Test Samples": 4430,
            "Precision": "76.58%",
            "Recall": "98.24%",
            "F1-Score": "86.07%"
        }
    ]
    df_all_cm = pd.DataFrame(all_cm)
    df_all_cm.to_excel(writer, sheet_name="All Models Error Comparison", index=False)

    # ---------------- 5. Methodology & Technical Architecture ----------------
    tech_data = [
        {"Parameter / Dimension": "Project Title", "Specification / Architecture": "Horus AI Amazon Market Sentiment & Product Intelligence"},
        {"Parameter / Dimension": "Academic Track", "Specification / Architecture": "Data Analysis & AI Graduation Project (Growth Level Academy)"},
        {"Parameter / Dimension": "Project Supervisor", "Specification / Architecture": "Eng. Aya Badwy"},
        {"Parameter / Dimension": "Winning Model Algorithm", "Specification / Architecture": "Calibrated Linear Support Vector Machine (LinearSVC with CalibratedClassifierCV)"},
        {"Parameter / Dimension": "Final Test Accuracy", "Specification / Architecture": "95.01% on 4,430 independent unseen test reviews"},
        {"Parameter / Dimension": "Final Test F1-Score", "Specification / Architecture": "95.36% (Harmonic Mean of Precision 95.26% & Recall 95.46%)"},
        {"Parameter / Dimension": "Cross-Validation Stability", "Specification / Architecture": "5-Fold Stratified Cross-Validation F1 Mean = 95.03% (Variance < 0.003)"},
        {"Parameter / Dimension": "Feature Extraction", "Specification / Architecture": "Bilingual TF-IDF Vectorizer with 12,000 N-Gram features (Unigrams & Bigrams)"},
        {"Parameter / Dimension": "Tokenization & Regex", "Specification / Architecture": "Unicode Word Boundary Regex `(?u)\\b\\w+\\b` supporting Arabic, English & French/German"},
        {"Parameter / Dimension": "Negation Handling", "Specification / Architecture": "Explicit retention of negative cues (not, never, no, مش, لم, لا, غير, بدون)"},
        {"Parameter / Dimension": "Probability Calibration", "Specification / Architecture": "Sigmoid/Platt Scaling via CalibratedClassifierCV for interpretable confidence scores"},
        {"Parameter / Dimension": "Inference Latency", "Specification / Architecture": "< 2 milliseconds per review on standard CPU"}
    ]
    df_tech = pd.DataFrame(tech_data)
    df_tech.to_excel(writer, sheet_name="Methodology & Architecture", index=False)

# ---------------- OpenPyXL Advanced Executive Styling ----------------
wb = openpyxl.load_workbook(excel_en_path)

header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Navy Blue
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

winner_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid") # Light Emerald Green
winner_font = Font(name="Segoe UI", size=10, bold=True, color="065F46")

regular_font = Font(name="Segoe UI", size=10, color="1E293B")
bold_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")

thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

pct_cols_bm = [3, 4, 5, 6, 7] # Accuracy, F1, Precision, Recall, CV Mean

for ws in wb.worksheets:
    ws.views.sheetView[0].rightToLeft = False # Standard English LTR
    
    # Header row formatting
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # Data rows formatting
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        ws.row_dimensions[row_idx].height = 22
        is_linear_svm = False
        if ws.title in ["Models Benchmark", "All Models Error Comparison"]:
            if "Linear SVM" in str(row[1].value):
                is_linear_svm = True

        for col_idx, cell in enumerate(row, start=1):
            cell.border = thin_border
            
            # Format percentages in Benchmark sheet
            if ws.title == "Models Benchmark" and col_idx in pct_cols_bm:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '0.00%'
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00'
                else:
                    cell.number_format = '#,##0'
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Apply winner highlighting
            if is_linear_svm:
                cell.fill = winner_fill
                cell.font = winner_font
            else:
                cell.font = regular_font

    # Auto-adjust column widths
    for col_idx, col in enumerate(ws.columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in col:
            val = cell.value
            if val is not None:
                if isinstance(val, float) and ws.title == "Models Benchmark" and col_idx in pct_cols_bm:
                    val_str = f"{val*100:.2f}%"
                else:
                    val_str = str(val)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

wb.save(excel_en_path)
print(f"Successfully generated executive English Excel report at: {excel_en_path}")
