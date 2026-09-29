import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_normal_excel(filepath):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.views.sheetView[0].rightToLeft = False
    ws.showGridLines = True

    # Standard headers starting directly at Row 1
    headers = [
        "Rank",
        "Model (Algorithm)",
        "Accuracy",
        "F1-Score",
        "Precision",
        "Recall",
        "5-Fold CV Mean",
        "Evaluation / Status"
    ]

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Classic Navy Blue
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # Standard rows data
    data = [
        [1, "Linear SVM (Calibrated)", 0.9501, 0.9536, 0.9526, 0.9546, 0.9503, "Production Champion (Best Overall Balance)"],
        [2, "Logistic Regression", 0.9481, 0.9519, 0.9479, 0.9559, 0.9481, "Strong Competitor (High Accuracy & Stable Probabilities)"],
        [3, "Multinomial Naive Bayes", 0.9354, 0.9416, 0.9155, 0.9693, 0.9407, "Fast Baseline (High Recall, Higher False Positives)"],
        [4, "Random Forest", 0.8291, 0.8607, 0.7658, 0.9824, 0.8624, "Biased towards Majority Class (Overfitting on TF-IDF)"]
    ]

    regular_font = Font(name="Segoe UI", size=10, color="0F172A")
    bold_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    winner_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
    winner_font = Font(name="Segoe UI", size=10, bold=True, color="166534")

    for r_idx, row_values in enumerate(data, start=2):
        ws.row_dimensions[r_idx].height = 24
        is_winner = (r_idx == 2)
        for c_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border
            
            if is_winner:
                cell.fill = winner_fill
                cell.font = winner_font
            else:
                cell.font = regular_font

            if c_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c_idx in [3, 4, 5, 6, 7]:
                cell.number_format = '0.00%'
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 8:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_l = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[col_letter].width = max(max_l + 4, 12)

    try:
        wb.save(filepath)
        print(f"Normal Excel saved to: {filepath}")
    except PermissionError:
        alt_path = filepath.replace(".xlsx", "_Clean.xlsx")
        wb.save(alt_path)
        print(f"Original locked. Saved to alternate path: {alt_path}")

if __name__ == "__main__":
    targets = [
        r"C:\Users\Basem\Desktop\Models_Benchmark.xlsx",
        r"D:\project_model\models_benchmark_evaluation_en.xlsx",
        r"D:\project_model\models_benchmark_report.xlsx"
    ]
    for target in targets:
        create_normal_excel(target)
