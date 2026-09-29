import json
import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

base_dir = os.path.dirname(os.path.abspath(__file__))
json_path = os.path.join(base_dir, "models", "model_comparison.json")
excel_path = os.path.join(base_dir, "models_benchmark_report.xlsx")

with open(json_path, "r", encoding="utf-8") as f:
    data = json.load(f)

# Create a Pandas Excel Writer using openpyxl engine
with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
    # 1. Models Comparison Sheet
    models_dict = data.get("all_models_comparison", {})
    comp_rows = []
    rank = 1
    # Sort by F1 score descending
    sorted_models = sorted(models_dict.items(), key=lambda x: x[1]['f1_score'], reverse=True)
    for name, metrics in sorted_models:
        is_winner = "الفائز المعتمد (Production Winner)" if name == data.get("best_model") else "نموذج تجريبي مقارن"
        comp_rows.append({
            "الترتيب": rank,
            "الخوارزمية (Algorithm)": name,
            "الدقة (Accuracy)": f"{metrics['accuracy']*100:.2f}%",
            "معامل الدقة (Precision)": f"{metrics['precision']*100:.2f}%",
            "معامل الاستدعاء (Recall)": f"{metrics['recall']*100:.2f}%",
            "الـ F1-Score": f"{metrics['f1_score']*100:.2f}%",
            "5-Fold CV Mean": f"{metrics['cv_f1_mean']*100:.2f}%",
            "الحالة والاعتماد": is_winner
        })
        rank += 1
    
    df_comp = pd.DataFrame(comp_rows)
    df_comp.to_excel(writer, sheet_name="Models Benchmark", index=False)

    # 2. Confusion Matrix for Winning Model
    cm = data.get("final_confusion_matrix", [[1937, 113], [108, 2272]])
    cm_df = pd.DataFrame({
        "التصنيف الفعلي \\ التوقع": ["مراجعة سلبية فعلية (Actual Negative)", "مراجعة إيجابية فعلية (Actual Positive)"],
        "توقع سلبي (Predicted Negative)": [cm[0][0], cm[1][0]],
        "توقع إيجابي (Predicted Positive)": [cm[0][1], cm[1][1]],
        "الإجمالي الفعلي": [cm[0][0] + cm[0][1], cm[1][0] + cm[1][1]]
    })
    cm_df.to_excel(writer, sheet_name="Confusion Matrix", index=False)

    # 3. All Models Confusion Matrices
    all_cm_rows = []
    for name, metrics in sorted_models:
        m_cm = metrics.get("confusion_matrix", [[0,0],[0,0]])
        all_cm_rows.append({
            "النموذج": name,
            "True Negative (سلبي صحيح)": m_cm[0][0],
            "False Positive (إيجابي خاطئ)": m_cm[0][1],
            "False Negative (سلبي خاطئ)": m_cm[1][0],
            "True Positive (إيجابي صحيح)": m_cm[1][1],
            "إجمالي العينة": sum(m_cm[0]) + sum(m_cm[1])
        })
    df_all_cm = pd.DataFrame(all_cm_rows)
    df_all_cm.to_excel(writer, sheet_name="All Confusion Matrices", index=False)

    # 4. Pipeline & Methodology Specs
    meta_rows = [
        {"المعيار / الخاصية": "المسار الأكاديمي", "التفاصيل والمواصفات": "مشروع تخرج الذكاء الاصطناعي وتحليل البيانات - برنامج حورس (Growth Level Academy)"},
        {"المعيار / الخاصية": "المشرف الأكاديمي", "التفاصيل والمواصفات": "Eng. Aya Badwy (المهندسة آية بدوي)"},
        {"المعيار / الخاصية": "النموذج الفائز (Winning Algorithm)", "التفاصيل والمواصفات": data.get("best_model", "Linear SVM (Calibrated)")},
        {"المعيار / الخاصية": "دقة الاختبار (Test Accuracy)", "التفاصيل والمواصفات": f"{data.get('final_accuracy', 0.9501)*100:.2f}%"},
        {"المعيار / الخاصية": "معامل F1 (Test F1-Score)", "التفاصيل والمواصفات": f"{data.get('final_f1', 0.9536)*100:.2f}%"},
        {"المعيار / الخاصية": "طريقة المعالجة (Feature Extraction)", "التفاصيل والمواصفات": "Bilingual TF-IDF Vectorizer (12,000 N-Gram Features, Unigrams + Bigrams)"},
        {"المعيار / الخاصية": "معالجة النصوص (NLP Preprocessing)", "التفاصيل والمواصفات": "تنظيف HTML والرموز، مع الاحتفاظ الصارم بأدوات النفي (not, never, no, مش، لم، لا)"},
        {"المعيار / الخاصية": "بروتوكول التحقق الإحصائي", "التفاصيل والمواصفات": "5-Fold Stratified Cross-Validation + 20% Holdout Test Set"},
        {"المعيار / الخاصية": "موازنة البيانات (Class Balancing)", "التفاصيل والمواصفات": "عينات متوازنة بين الإيجابي والسلبي لمنع التحيز وضمان التكافؤ"},
        {"المعيار / الخاصية": "معايرة الاحتماليات (Probability Calibration)", "التفاصيل والمواصفات": "CalibratedClassifierCV (Isotonic/Sigmoid) لإنتاج نسب ثقة حقيقية قابلة للتفسير"},
        {"المعيار / الخاصية": "حفظ النماذج (Model Persistence)", "التفاصيل والمواصفات": "Joblib Serialized Binary (sentiment_model.pkl & tfidf_vectorizer.pkl)"}
    ]
    df_meta = pd.DataFrame(meta_rows)
    df_meta.to_excel(writer, sheet_name="Methodology & Architecture", index=False)

# Styling with openpyxl
wb = openpyxl.load_workbook(excel_path)

header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
winner_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
regular_font = Font(name="Segoe UI", size=10)
thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

for ws in wb.worksheets:
    ws.views.sheetView[0].rightToLeft = True
    for col_idx, col in enumerate(ws.columns, 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx, cell in enumerate(col, 1):
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", horizontal="center" if row_idx > 1 and col_idx in [1,3,4,5,6,7] else "right")
            if row_idx == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif ws.title == "Models Benchmark" and row_idx == 2:
                # Highlight winner
                cell.fill = winner_fill
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="065F46")
            
            val_str = str(cell.value or '')
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)
    ws.row_dimensions[1].height = 28

wb.save(excel_path)
print(f"Excel created successfully at {excel_path} ({os.path.getsize(excel_path)} bytes)")
