import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_excel(output_path):
    wb = openpyxl.Workbook()
    
    # ---------------- SHEET 1: Models Benchmark ----------------
    ws1 = wb.active
    ws1.title = "Models Benchmark"
    ws1.views.sheetView[0].rightToLeft = False
    ws1.showGridLines = True

    # Title matching the screenshot
    ws1['A1'] = "Sheet 1: Models Benchmark"
    ws1['A1'].font = Font(name="Segoe UI", size=12, bold=True, color="DC2626") # Red color like screenshot
    ws1['A1'].alignment = Alignment(vertical="center")
    ws1.row_dimensions[1].height = 25
    ws1.row_dimensions[2].height = 10 # Spacer row

    # Table Headers at Row 3
    headers = [
        "Rank",
        "Model\n(Algorithm)",
        "Accuracy",
        "F1-\nScore",
        "Precision",
        "Recall",
        "5-Fold\nCV Mean",
        "Evaluation / Status"
    ]
    
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid") # Modern soft gray/slate
    header_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    
    border_color = "CBD5E1" # Crisp clean light slate border
    table_border = Border(
        left=Side(style='thin', color=border_color),
        right=Side(style='thin', color=border_color),
        top=Side(style='thin', color=border_color),
        bottom=Side(style='thin', color=border_color)
    )

    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws1.cell(row=3, column=col_idx, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = table_border
    ws1.row_dimensions[3].height = 34

    # Rows Data
    rows_data = [
        (
            1,
            "Linear SVM\n(Calibrated)",
            0.9501,
            0.9536,
            0.9526,
            0.9546,
            0.9503,
            "🏆 Production\nChampion (Best Overall\nBalance)"
        ),
        (
            2,
            "Logistic\nRegression",
            0.9481,
            0.9519,
            0.9479,
            0.9559,
            0.9481,
            "Strong Competitor\n(High Accuracy & Stable\nProbabilities)"
        ),
        (
            3,
            "Multinomial\nNaive Bayes",
            0.9354,
            0.9416,
            0.9155,
            0.9693,
            0.9407,
            "Fast Baseline (High\nRecall, Higher False\nPositives)"
        ),
        (
            4,
            "Random Forest",
            0.8291,
            0.8607,
            0.7658,
            0.9824,
            0.8624,
            "Biased towards Majority\nClass (Overfitting on TF-\nIDF)"
        )
    ]

    regular_font = Font(name="Segoe UI", size=10, color="0F172A")
    bold_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")

    for r_offset, r_values in enumerate(rows_data, start=4):
        ws1.row_dimensions[r_offset].height = 42
        for c_idx, val in enumerate(r_values, start=1):
            cell = ws1.cell(row=r_offset, column=c_idx, value=val)
            cell.border = table_border
            
            # Rank column
            if c_idx == 1:
                cell.font = bold_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            # Model column
            elif c_idx == 2:
                cell.font = bold_font if r_offset == 4 else regular_font
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            # Metric percentages
            elif c_idx in [3, 4, 5, 6, 7]:
                cell.number_format = '0.00%'
                cell.font = bold_font if r_offset == 4 else regular_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            # Status column
            elif c_idx == 8:
                cell.font = bold_font if r_offset == 4 else regular_font
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Specific clean column widths
    col_widths = {
        'A': 8,   # Rank
        'B': 18,  # Model
        'C': 13,  # Accuracy
        'D': 12,  # F1-Score
        'E': 13,  # Precision
        'F': 12,  # Recall
        'G': 14,  # 5-Fold CV Mean
        'H': 32   # Evaluation / Status
    }
    for col_letter, width in col_widths.items():
        ws1.column_dimensions[col_letter].width = width

    # ---------------- SHEET 2: Train-Test Split Specs ----------------
    ws2 = wb.create_sheet(title="Train-Test Split Specs")
    ws2.views.sheetView[0].rightToLeft = False
    ws2.showGridLines = True

    ws2['A1'] = "Dataset Splitting & Training Methodology"
    ws2['A1'].font = Font(name="Segoe UI", size=12, bold=True, color="1E3A8A")
    ws2.row_dimensions[1].height = 25

    split_headers = ["Dataset Split", "Percentage", "Samples Count", "Sampling Strategy", "Academic Purpose"]
    for col_idx, h_text in enumerate(split_headers, start=1):
        cell = ws2.cell(row=3, column=col_idx, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = table_border
    ws2.row_dimensions[3].height = 28

    split_rows = [
        ("Training Set", 0.80, 17720, "Stratified Random (random_state=42)", "Algorithm learning & TF-IDF vocabulary fitting"),
        ("Testing Set (Holdout)", 0.20, 4430, "Stratified Holdout (Unseen)", "Unbiased final model benchmarking"),
        ("Total Corpus", 1.00, 22150, "Bilingual Augmented (Arabic + English)", "Multi-dialect, multi-category coverage"),
        ("5-Fold Cross Validation", 0.20, 3544, "Stratified 5-Fold per split", "Ensures no overfitting across folds"),
        ("Positive Class (4 & 5 Stars)", 0.538, 11917, "Balanced Binary Mapping", "Representative sample of satisfied customers"),
        ("Negative Class (1 & 2 Stars)", 0.462, 10233, "Balanced Binary Mapping", "Representative sample of complaints & defects")
    ]

    for r_idx, r_data in enumerate(split_rows, start=4):
        ws2.row_dimensions[r_idx].height = 24
        for c_idx, val in enumerate(r_data, start=1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = table_border
            if c_idx == 2:
                cell.number_format = '0.0%'
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 3:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws2.columns:
        col_letter = get_column_letter(col[0].column)
        max_l = max(len(str(c.value or '')) for c in col)
        ws2.column_dimensions[col_letter].width = max(max_l + 4, 14)

    # ---------------- SHEET 3: Confusion Matrix ----------------
    ws3 = wb.create_sheet(title="Confusion Matrix")
    ws3.views.sheetView[0].rightToLeft = False
    ws3.showGridLines = True

    ws3['A1'] = "Confusion Matrix: Linear SVM (Calibrated) on Holdout Test Set (4,430 Samples)"
    ws3['A1'].font = Font(name="Segoe UI", size=12, bold=True, color="1E3A8A")
    ws3.row_dimensions[1].height = 25

    cm_headers = ["Actual Class \\ Predicted Class", "Predicted Negative", "Predicted Positive", "Actual Total", "Class Accuracy"]
    for col_idx, h_text in enumerate(cm_headers, start=1):
        cell = ws3.cell(row=3, column=col_idx, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = table_border
    ws3.row_dimensions[3].height = 28

    cm_data = [
        ("Actual Negative (Complaints / Defects)", 1937, 113, 2050, "94.49% (Specificity)"),
        ("Actual Positive (Satisfied Customers)", 108, 2272, 2380, "95.46% (Sensitivity)"),
        ("Total Predicted", 2045, 2385, 4430, "95.01% (Overall Accuracy)")
    ]

    for r_idx, r_data in enumerate(cm_data, start=4):
        ws3.row_dimensions[r_idx].height = 24
        for c_idx, val in enumerate(r_data, start=1):
            cell = ws3.cell(row=r_idx, column=c_idx, value=val)
            cell.font = bold_font if r_idx == 6 else regular_font
            cell.border = table_border
            if c_idx in [2, 3, 4]:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 5:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws3.columns:
        col_letter = get_column_letter(col[0].column)
        max_l = max(len(str(c.value or '')) for c in col)
        ws3.column_dimensions[col_letter].width = max(max_l + 4, 16)

    # Save workbook
    wb.save(output_path)
    print(f"Saved: {output_path}")

if __name__ == "__main__":
    targets = [
        r"D:\project_model\models_benchmark_evaluation_en.xlsx",
        r"D:\project_model\models_benchmark_report.xlsx",
        r"C:\Users\Basem\Desktop\Models_Benchmark.xlsx"
    ]
    for target in targets:
        try:
            build_excel(target)
        except Exception as e:
            print(f"Error saving to {target}: {e}")
